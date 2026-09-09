"""辅助功能的真实 HTTP/文件输出与微信签名失败、并发回归；上游全部隔离。"""
import hashlib
import io
import json
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest import mock
from xml.etree import ElementTree

from openpyxl import load_workbook
from tests.test_security_reliability import server, TestClient, make_request


class PublicFeatureTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(server.app)
        self.addCleanup(self.client.close)
        self.data = {
            'item_id': '7880000000000000001', 'kind': 'video', 'platform': 'douyin',
            'title': '回归视频', 'author': '测试作者',
            'stats': {'digg': 0, 'comment': 12, 'collect': 5, 'share': 8},
            'duration_ms': 10250, 'video': {'width': 1080, 'height': 1920},
        }

    def share(self):
        share = server._share_create(make_request(), self.data, '')
        self.addCleanup(server.db_exec, 'DELETE FROM shares WHERE id=?', (share['sid'],))
        return share['sid']

    def test_export_is_real_workbook_with_title_zero_and_all_statistics(self):
        response = self.client.post('/api/export/xlsx', json={'items': [self.data]})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'PK'))
        self.assertIn('attachment', response.headers['content-disposition'])
        book = load_workbook(io.BytesIO(response.content))
        self.addCleanup(book.close)
        sheet = book.active
        row = list(sheet.values)[1]
        self.assertEqual(row[2:5], ('回归视频', '测试作者', self.data['item_id']))
        self.assertEqual(row[5:11], (10.2, '1080×1920', 0, 12, 5, 8))
        self.assertEqual(sheet.freeze_panes, 'A2')

    def test_export_formula_injection_is_inert_in_saved_file(self):
        item = dict(self.data, title='=HYPERLINK("https://example.test")',
                    author='+1+1', tags=['@demo'], _link='\t=1+1')
        response = self.client.post('/api/export/xlsx', json={'items': [item]})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content), data_only=False)
        self.addCleanup(book.close)
        for column in ('C', 'D', 'R'):
            cell = book.active[column + '2']
            self.assertEqual(cell.data_type, 's')
            self.assertTrue(cell.value.startswith("'"))

    def test_export_note_and_malformed_optional_metadata(self):
        item = {'kind': 'note', 'images': [{'url': 'https://example.test/a.jpg'}, None],
                'video': None, 'stats': [], 'duration_ms': 'NaN', 'create_time': 'bad'}
        response = self.client.post('/api/export/xlsx', json={'items': [item]})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.addCleanup(book.close)
        row = list(book.active.values)[1]
        self.assertEqual(row[1], '图集')
        self.assertIsNone(row[5])
        self.assertEqual(row[15], 'https://example.test/a.jpg')

    def test_export_rejects_oversized_item_count(self):
        response = self.client.post('/api/export/xlsx',
                                    json={'items': [{}] * (server.EXPORT_ITEMS_MAX + 1)})
        self.assertEqual(response.status_code, 422)

    def test_qr_png_and_svg_are_real_images_without_upstream_calls(self):
        sid = self.share()
        with mock.patch.object(server, '_parse_share', side_effect=AssertionError('不能重新解析')):
            png = self.client.get(f'/s/{sid}/qr.png')
            svg = self.client.get(f'/s/{sid}/qr.svg')
        self.assertEqual(png.status_code, 200)
        self.assertTrue(png.content.startswith(b'\x89PNG\r\n\x1a\n'))
        self.assertGreater(len(png.content), 100)
        self.assertEqual(svg.status_code, 200)
        self.assertEqual(ElementTree.fromstring(svg.content).tag,
                         '{http://www.w3.org/2000/svg}svg')

    def test_qr_keeps_assigned_origin_when_primary_changes(self):
        import segno
        sid = self.share()
        server.db_exec('UPDATE shares SET assigned_origin=? WHERE id=?',
                       ('https://share.example.test', sid))
        with mock.patch.object(segno, 'make', wraps=segno.make) as encode:
            self.assertEqual(self.client.get(f'/s/{sid}/qr.png').status_code, 200)
        self.assertEqual(encode.call_args.args[0], f'https://share.example.test/s/{sid}')

    def test_report_deduplicates_and_admin_can_handle_it(self):
        sid = self.share()
        self.addCleanup(server.db_exec, 'DELETE FROM reports WHERE sid=?', (sid,))
        payload = {'sid': sid, 'reason': '测试投诉', 'contact': 'test@example.test'}
        with mock.patch.dict(server._report_hits, {}, clear=True):
            for _ in range(2):
                self.assertEqual(self.client.post('/api/report', json=payload).status_code, 200)
        rows = server.db_exec('SELECT * FROM reports WHERE sid=?', (sid,), 'all')
        self.assertEqual(len(rows), 1)
        self.assertNotEqual(rows[0]['ip'], 'testclient')
        with mock.patch.object(server, '_require_admin'):
            response = self.client.post(f"/api/admin/reports/{rows[0]['id']}/handle")
            report = self.client.get('/api/admin/reports').json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(next(r for r in report['reports'] if r['sid'] == sid)['handled'], 1)

    def test_report_missing_and_expired_share_rejected(self):
        sid = self.share()
        server.db_exec('UPDATE shares SET expires_at=1 WHERE id=?', (sid,))
        with mock.patch.dict(server._report_hits, {}, clear=True):
            for target in (sid, 'missingReviewShare'):
                response = self.client.post('/api/report', json={'sid': target, 'reason': '测试'})
                self.assertEqual(response.status_code, 404)

    def test_auxiliary_admin_routes_require_authentication(self):
        for method, path, body in (
            ('GET', '/api/admin/share-config', None),
            ('POST', '/api/admin/share-config', {'wx_appid': 'attacker'}),
            ('GET', '/api/admin/reports', None),
            ('POST', '/api/admin/reports/1/handle', {}),
            ('GET', '/api/admin/traffic-stats', None),
            ('GET', '/api/admin/play-stats', None),
            ('GET', '/api/admin/play-logs', None),
            ('GET', '/api/admin/mihomo', None),
            ('POST', '/api/admin/mihomo', {'action': 'stop'}),
        ):
            with self.subTest(path=path, method=method):
                self.assertIn(self.client.request(method, path, json=body).status_code, (401, 403))

    def test_traffic_read_flushes_once_and_separates_download_from_play(self):
        server.db_exec('DELETE FROM media_traffic')
        self.addCleanup(server.db_exec, 'DELETE FROM media_traffic')
        with mock.patch.dict(server._traffic_pending, {}, clear=True), \
                mock.patch.object(server, '_require_admin'):
            server._traffic_add('download', 1024)
            server._traffic_add('play', 512)
            first = self.client.get('/api/admin/traffic-stats?days=1').json()
            second = self.client.get('/api/admin/traffic-stats?days=1').json()
        self.assertEqual(first, second)
        self.assertEqual(first['totals'], {'play_requests': 1, 'play_bytes': 512,
                                         'download_requests': 1, 'download_bytes': 1024})

    def test_public_pages_have_version_and_no_unexpanded_placeholders(self):
        for path in ('/', '/transcript', '/api-docs', '/api-console', '/admin_d'):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers['x-app-version'], server.APP_VERSION)
                self.assertIn('no-store', response.headers['cache-control'])
                self.assertNotIn('{{FRONTEND_VERSION}}', response.text)
                self.assertNotIn('{{SEO_HEAD}}', response.text)
        self.assertEqual(self.client.get('/healthz').json()['version'], server.APP_VERSION)

    def test_seo_and_og_resources_exist(self):
        for path, expected in (('/robots.txt', 'Sitemap:'), ('/llms.txt', '#'),
                               ('/sitemap.xml', 'urlset'), ('/og.svg', '<svg')):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected, response.text)
        self.assertTrue(self.client.get('/og.png').content.startswith(b'\x89PNG'))


class WechatSignatureTests(unittest.TestCase):
    keys = ('wx_appid', 'wx_secret', 'wx_ticket', 'wx_ticket_exp', 'share_primary_domain')

    def setUp(self):
        self.previous = {key: server.app_setting(key) for key in self.keys}
        self.addCleanup(self.restore_settings)
        for key in self.keys:
            server.set_app_setting(key, '')
        self.client = TestClient(server.app)
        self.addCleanup(self.client.close)
        self.request = make_request()
        self.page = 'http://testserver/s/demo?lang=en'

    def restore_settings(self):
        for key, value in self.previous.items():
            server.set_app_setting(key, value)

    def configure(self):
        server.set_app_setting('wx_appid', 'test-app')
        server.set_app_setting('wx_secret', 'private-test-secret')

    def fetch(self):
        return self.client.get('/api/wx/jssdk', params={'url': self.page})

    def test_unconfigured_is_disabled_without_network(self):
        with mock.patch.object(server, '_wx_api', side_effect=AssertionError('不能访问上游')):
            self.assertEqual(self.fetch().json(), {'enabled': False})

    def test_external_origin_rejected_before_fetch(self):
        self.configure()
        with mock.patch.object(server, '_wx_api') as upstream:
            response = self.client.get('/api/wx/jssdk', params={'url': 'https://evil.test/s/demo'})
        self.assertEqual(response.status_code, 403)
        upstream.assert_not_called()

    def test_signature_uses_exact_url_without_fragment_and_cached_ticket(self):
        self.configure()
        with mock.patch.object(server, '_wx_api', side_effect=[
                {'access_token': 'private-token'}, {'ticket': 'private-ticket', 'expires_in': 7200}]) as upstream:
            first = server.wx_jssdk(self.request, self.page + '#fragment')
            second = server.wx_jssdk(self.request, self.page)
        self.assertTrue(first['enabled'])
        self.assertTrue(second['enabled'])
        self.assertEqual(upstream.call_count, 2)
        raw = (f"jsapi_ticket=private-ticket&noncestr={first['nonceStr']}"
               f"&timestamp={first['timestamp']}&url={self.page}")
        self.assertEqual(first['signature'], hashlib.sha1(raw.encode()).hexdigest())
        self.assertNotIn('private-', json.dumps(first))

    def test_failures_are_neutral_disabled_json_and_do_not_cache(self):
        self.configure()
        cases = (
            [RuntimeError('private-test-secret in URL')],
            [{'errcode': 401, 'errmsg': 'private-test-secret'}],
            [None],
            [{'access_token': 'private-token'}, {'errmsg': 'private-token'}],
            [{'access_token': 'private-token'}, {'ticket': 'private-ticket', 'expires_in': 'bad'}],
            [{'access_token': 'private-token'}, {'ticket': 'private-ticket', 'expires_in': -1}],
        )
        for values in cases:
            with self.subTest(values=values), mock.patch.object(server, '_wx_api', side_effect=values):
                response = self.fetch()
                self.assertEqual(response.status_code, 200)
                self.assertFalse(response.json()['enabled'])
                self.assertNotIn('private-', response.text)
                self.assertEqual(server.app_setting('wx_ticket'), '')

    def test_concurrent_requests_share_one_refresh(self):
        self.configure()
        def upstream(url):
            time.sleep(0.01)
            return ({'access_token': 'token'} if '/token?' in url
                    else {'ticket': 'ticket', 'expires_in': 7200})
        with mock.patch.object(server, '_wx_api', side_effect=upstream) as fetch:
            with ThreadPoolExecutor(max_workers=8) as pool:
                results = list(pool.map(lambda _: server._wx_ticket(), range(8)))
        self.assertEqual(fetch.call_count, 2)
        self.assertEqual(results, [('test-app', 'ticket')] * 8)

    def test_changing_either_credential_invalidates_ticket_but_empty_secret_preserves_it(self):
        self.configure()
        with mock.patch.object(server, '_require_admin'):
            for body in ({'wx_appid': 'next-app'}, {'wx_secret': 'next-secret'}):
                server.set_app_setting('wx_ticket', 'old-ticket')
                server.set_app_setting('wx_ticket_exp', time.time() + 5000)
                response = self.client.post('/api/admin/share-config', json=body)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(server.app_setting('wx_ticket'), '')
                self.assertNotIn('next-secret', response.text)
            server.set_app_setting('wx_ticket', 'valid-ticket')
            response = self.client.post('/api/admin/share-config', json={'wx_secret': ''})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(server.app_setting('wx_secret'), 'next-secret')
        self.assertEqual(server.app_setting('wx_ticket'), 'valid-ticket')

    def test_domain_normalization_and_invalid_origin_leave_current_domain(self):
        with mock.patch.object(server, '_require_admin'):
            response = self.client.post('/api/admin/share-config',
                                        json={'primary_domain': 'share.example.test/'})
            self.assertEqual(response.json()['primary_domain'], 'https://share.example.test')
            for value in ('https://user:pass@share.example.test', 'https://share.example.test/path'):
                response = self.client.post('/api/admin/share-config', json={'primary_domain': value})
                self.assertEqual(response.status_code, 422)
        self.assertEqual(server.app_setting('share_primary_domain'), 'https://share.example.test')
